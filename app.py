import re
import pandas as pd
import streamlit as st
import matplotlib.pyplot as plt

from pypdf import PdfReader
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


st.set_page_config(
    page_title="AI Resume Analyzer",
    page_icon="📄",
    layout="wide"
)

st.title("📄 AI Resume Analyzer & Job Recommendation System")
st.write(
    "Upload your resume to identify skills, explore suitable job roles, "
    "and discover areas for improvement."
)


JOB_ROLES = {
    "Data Analyst": [
        "Python", "SQL", "Excel", "Pandas", "NumPy", "Power BI"
    ],
    "Machine Learning Engineer": [
        "Python", "Machine Learning", "Scikit-learn",
        "Pandas", "NumPy", "FastAPI", "Docker"
    ],
    "AI Engineer": [
        "Python", "Deep Learning", "NLP", "TensorFlow",
        "PyTorch", "APIs", "LLM"
    ],
    "Web Developer": [
        "HTML", "CSS", "JavaScript", "React",
        "Node.js", "SQL", "Git"
    ],
    "NLP Engineer": [
        "Python", "NLP", "Transformers",
        "Hugging Face", "Machine Learning", "PyTorch"
    ],
    "Computer Vision Engineer": [
        "Python", "OpenCV", "CNN", "Deep Learning",
        "TensorFlow", "PyTorch"
    ]
}

SKILLS = [
    "Python", "Java", "C", "C++", "JavaScript", "HTML", "CSS",
    "R", "SQL", "MySQL", "PostgreSQL", "MongoDB", "Pandas",
    "NumPy", "Excel", "Power BI", "Matplotlib", "Statistics",
    "Data Analysis", "Machine Learning", "Deep Learning",
    "Scikit-learn", "TensorFlow", "PyTorch", "NLP", "LLM",
    "Transformers", "Hugging Face", "OpenCV", "CNN", "React",
    "Node.js", "Flask", "FastAPI", "REST API", "Docker",
    "Git", "GitHub", "AWS", "Azure", "Streamlit", "Linux"
]


def extract_text(uploaded_file):
    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        return "\n".join(
            page.extract_text() or "" for page in reader.pages
        )

    if filename.endswith(".docx"):
        document = Document(uploaded_file)
        parts = [
            paragraph.text for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)

        return "\n".join(parts)

    raise ValueError("Please upload a PDF or DOCX resume.")


def clean_text(text):
    text = text.lower()
    text = text.replace("scikit learn", "scikit-learn")
    text = text.replace("node js", "node.js")
    text = re.sub(r"\S+@\S+", " ", text)
    text = re.sub(r"https?://\S+|www\.\S+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def find_skills(text):
    detected = []

    for skill in SKILLS:
        pattern = r"(?<!\w)" + re.escape(skill.lower()) + r"(?!\w)"
        if re.search(pattern, text):
            detected.append(skill)

    return sorted(set(detected), key=str.lower)


def analyze_resume(text, detected_skills):
    roles = list(JOB_ROLES.keys())
    role_texts = [" ".join(JOB_ROLES[role]) for role in roles]

    vectorizer = TfidfVectorizer(
        lowercase=True, ngram_range=(1, 2)
    )
    matrix = vectorizer.fit_transform([text] + role_texts)

    similarities = cosine_similarity(
        matrix[0:1], matrix[1:]
    ).flatten()

    detected_set = {skill.lower() for skill in detected_skills}
    results = []

    for i, role in enumerate(roles):
        required = JOB_ROLES[role]
        matched = [
            skill for skill in required
            if skill.lower() in detected_set
        ]
        missing = [
            skill for skill in required
            if skill.lower() not in detected_set
        ]

        coverage = len(matched) / len(required) * 100

        results.append({
            "Job Role": role,
            "Skill Coverage (%)": round(coverage, 2),
            "Text Similarity (%)": round(similarities[i] * 100, 2),
            "Matched Skills": ", ".join(matched) or "None detected",
            "Missing Skills": ", ".join(missing) or "None"
        })

    return pd.DataFrame(results).sort_values(
        ["Skill Coverage (%)", "Text Similarity (%)"],
        ascending=False
    ).reset_index(drop=True)


uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx"]
)

if uploaded_file is not None:
    try:
        raw_text = extract_text(uploaded_file)

        if not raw_text.strip():
            st.error(
                "No readable text was found. Try a text-based PDF or DOCX."
            )
        else:
            cleaned = clean_text(raw_text)
            skills = find_skills(cleaned)
            results = analyze_resume(cleaned, skills)

            st.success("Resume analyzed successfully!")

            st.subheader("1. Skills detected")
            if skills:
                st.write(", ".join(skills))
            else:
                st.info("No skills from the current dictionary were detected.")

            st.subheader("2. Top 3 job recommendations")
            top_three = results.head(3)
            st.dataframe(top_three, use_container_width=True)

            st.subheader("3. Skill coverage by job role")
            chart_data = results.set_index("Job Role")[
                "Skill Coverage (%)"
            ]
            st.bar_chart(chart_data)

            st.subheader("4. Personalized skill-gap analysis")
            target_role = st.selectbox(
                "Choose your target role",
                list(JOB_ROLES.keys())
            )

            selected = results[
                results["Job Role"] == target_role
            ].iloc[0]

            col1, col2 = st.columns(2)

            with col1:
                st.write("**Skills found**")
                st.write(selected["Matched Skills"])

            with col2:
                st.write("**Skills to learn**")
                missing = selected["Missing Skills"]

                if missing == "None":
                    st.write("No listed skill gaps detected.")
                else:
                    for skill in missing.split(", "):
                        st.write(f"- {skill}")

            st.caption(
                "Recommendations are estimates based on listed skills and "
                "text similarity. They do not determine actual job eligibility."
            )

    except Exception as error:
        st.error(f"Could not analyze the resume: {error}")
